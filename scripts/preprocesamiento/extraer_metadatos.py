#python -m preprocesamiento.extraer_metadatos

import os
import re
import json
import sys
import fitz
import docx
from pathlib import Path
from typing import Optional, List, Literal
from pydantic import BaseModel, Field
from langchain_openai import AzureChatOpenAI
from langchain_core.prompts import ChatPromptTemplate
from langchain_community.callbacks import get_openai_callback
from scripts.modelo.esquemas_pydantic import FiltrosMetadatos
from dotenv import load_dotenv


def limpiar_texto_basico(texto):
    return " ".join(texto.split())

def extraer_todo_el_documento(ruta):
    texto_completo = ""
    if ruta.lower().endswith(".pdf"):
        with fitz.open(ruta) as doc:
            for pagina in doc:
                texto_completo += pagina.get_text("text") + "\n"
    elif ruta.lower().endswith(".docx"):
        doc_word = docx.Document(ruta)
        for p in doc_word.paragraphs:
            if p.text.strip():
                texto_completo += p.text + "\n"
        for tabla in doc_word.tables:
            for fila in tabla.rows:
                fila_texto = " | ".join([c.text.replace("\n", " ").strip() for c in fila.cells if c.text.strip()])
                if fila_texto:
                    texto_completo += fila_texto + "\n"
    return limpiar_texto_basico(texto_completo)

def procesar_documento(ruta, llm):
    nombre = os.path.basename(ruta)
    print(f"\nProcesando: {nombre}")
    
    texto_documento = extraer_todo_el_documento(str(ruta))
    datos_finales = {"archivo": nombre}
    coste_peticion = 0.0
    
    prompt = ChatPromptTemplate.from_template("""
    Eres un auditor experto en contratacion publica de España.
    Tu tarea es analizar el siguiente Pliego de Clausulas Administrativas Particulares (PCAP).
    Extrae todos los metadatos solicitados basandote EXCLUSIVAMENTE en el texto proporcionado.
    Si un dato no aparece claramente especificado en el texto, dejalo nulo.
    
    TEXTO DEL PLIEGO:
    '''{texto}'''
    """)
    
    llm_estructurado = llm.with_structured_output(FiltrosMetadatos)
    cadena = prompt | llm_estructurado
    
    try:
        with get_openai_callback() as cb:
            resultado_pydantic = cadena.invoke({"texto": texto_documento})
            
            # Calculo del coste de esta peticion especifica
            coste_peticion = ((cb.prompt_tokens/1000000)*1.09) + ((cb.completion_tokens/1000000)*8.69)
            print(f"  -> Tokens procesados: {cb.total_tokens} | Coste doc: {coste_peticion:.4f} EUR")
            
            if resultado_pydantic:
                # Usamos model_dump para evitar el warning de obsolescencia
                datos_extrados = resultado_pydantic.model_dump(exclude_none=False)
                datos_finales.update(datos_extrados)
                
    except Exception as e:
        print(f"[ERROR] Procesando {nombre}: {e}")

    return datos_finales, coste_peticion

if __name__ == "__main__":

    
    
    # 1. Ajuste de rutas
    DIRECTORIO_ACTUAL = Path(__file__).resolve().parent
    DIRECTORIO_SCRIPTS = DIRECTORIO_ACTUAL.parent
    sys.path.append(str(DIRECTORIO_SCRIPTS))
    
    from modelo.esquemas_pydantic import FiltrosMetadatos

    # 2. Inicializacion exacta como en tu script original
    load_dotenv()
    
    api_key = os.environ.get("AZURE_OPENAI_API_KEY")
    endpoint = os.environ.get("AZURE_OPENAI_ENDPOINT")
    api_version = os.environ.get("AZURE_OPENAI_API_VERSION")
    
    if not all([api_key, endpoint, api_version]):
        raise ValueError("Faltan credenciales de Azure en el archivo .env. Revisa la configuracion.")
        
    print("Inicializando modelo Azure LLM...")
    llm_global = AzureChatOpenAI(
        azure_deployment="gpt-5.1",
        model_name="gpt-5.1",
        api_version=api_version,
        azure_endpoint=endpoint,
        api_key=api_key,
        temperature=0.0,
        max_tokens=16384
    )
    
    # 3. Configurar rutas
    BASE_DIR = Path(__file__).resolve().parent.parent          
    PROJECT_ROOT = BASE_DIR.parent           
    DIR_PLIEGOS = PROJECT_ROOT / "datos" / "pdfs" / "pcap"
    DIR_JSON = PROJECT_ROOT / "datos" / "jsons" / "metadatos.jsonl"
    
    if not DIR_PLIEGOS.exists():
        print(f"[ERROR] La carpeta {DIR_PLIEGOS} no existe.")
        exit(1)
        
    archivos_doc = list(DIR_PLIEGOS.glob("*.pdf")) + list(DIR_PLIEGOS.glob("*.docx"))
    DIR_JSON.parent.mkdir(parents=True, exist_ok=True)
    
    print(f"--- Iniciando Extraccion Masiva ({len(archivos_doc)} documentos) ---")
    
    coste_total_lote = 0.0

    with open(DIR_JSON, "w", encoding="utf-8") as metadatos_finales:
        for archivo in archivos_doc:
            diccionario_metadatos, coste_doc = procesar_documento(archivo, llm_global)
            
            coste_total_lote += coste_doc
            
            json_metadatos = json.dumps(diccionario_metadatos, ensure_ascii=False)
            metadatos_finales.write(json_metadatos + "\n")

    print("\n" + "="*50)
    print(" REPORTE FINAL DE COSTES DEL LOTE ")
    print(f"Documentos procesados: {len(archivos_doc)}")
    print(f"Coste total estimado:  {coste_total_lote:.4f} EUR")
    print("="*50)
    
    print(f"\n--- Proceso finalizado ---")
    print(f"Guardado en: {DIR_JSON}")