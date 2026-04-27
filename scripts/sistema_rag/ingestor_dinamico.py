# RECUERDA QUE ESTE SISTEMA ESTÁ BASADO EN QUE TÚ LE PONDRÁS EL NOMBRE A CADA ARCHIVO QUE SE SUBA
# TENDREMOS QUE GUARDAR EN UNA VARIABLE GLOBAL EL NÚMERO DE DOCUMENTOS QUE TENEMOS EN LA 4
# BASE DE DATOS VECTORIAL INICIAL Y AUMENTAR EN 1 CADA VEZ QUE SE SUBA UNO NUEVO
# ADEMÁS DE QUE PARA RENOMBRAR LOS ARCHIVOS DEBEREMOS USAR ESTA VARIABLE GLOBAL
# ASÍ TAMBIÉN MEJORAMOS LA SEGURIDAD
import fitz
from langchain_core.documents import Document
import docx

from scripts.preprocesamiento.extraer_metadatos import procesar_documento
from scripts.preprocesamiento.chunking_embeddings import(
    configurar_chunkeador, generar_id_hash, limpiar_texto, extraer_id)

def carga_diccionario_archivo_unico(dicc_metadatos):
    if not dicc_metadatos:
        print("El diccionario de metadatos esta vacio.")
        return {}, None, None

    nombre_archivo = dicc_metadatos.get("archivo")
    id_actual = None
    
    if nombre_archivo:
        id_actual = extraer_id(nombre_archivo)
        if not id_actual:
            print(f"Aviso: No se pudo extraer ID de {nombre_archivo}")
            
    return dicc_metadatos, nombre_archivo, id_actual


def ingestar_documento_individual(ruta_archivo, vector_db, llm, origen_tipo):
    """
    Procesa un documento (PDF o DOCX) y lo añade a la base de datos vectorial existente.
    """
    print(f"\n[INGESTA INCREMENTAL] Iniciando proceso para: {ruta_archivo}")
    
    # 1. Extraccion de metadatos con tu Pydantic + LLM
    print("[1/4] Extrayendo metadatos con IA...")
    
    resultado_procesamiento = procesar_documento(ruta_archivo, llm)
    
    # --- CORRECCION DEL BUG (Desempaquetado seguro) ---
    # Si la funcion devolvio una tupla (diccionario, coste, etc), nos quedamos solo con el diccionario (indice 0)
    if isinstance(resultado_procesamiento, tuple):
        dicc_metadatos = resultado_procesamiento[0]
    else:
        # Si devuelve solo el diccionario, lo dejamos tal cual
        dicc_metadatos = resultado_procesamiento
    # --------------------------------------------------
        
    diccionario_json, archivo, id_actual = carga_diccionario_archivo_unico(dicc_metadatos)
    
    # 2. Lectura del texto del Archivo
    print("[2/4] Leyendo contenido del documento...")
    docs_lista = [] 
    
    try:
        # LOGICA PARA PDF
        if ruta_archivo.lower().endswith(".pdf"):
            with fitz.open(ruta_archivo) as pdf_doc:
                for pagina in pdf_doc:
                    texto_limpio = limpiar_texto(pagina.get_text("text"))
                    if not texto_limpio:
                        continue   
                        
                    metadatos_final = {
                        "fuente": archivo,
                        "doc_id": id_actual if id_actual else "unknown",
                        "tipo_documento": "Archivo Externo (PDF)",
                        "pagina": pagina.number + 1,
                        "origen": origen_tipo 
                    }
                    
                    # Añadir metadatos json
                    tipos_permitidos = (str, int, float, bool)
                    for clave, valor in diccionario_json.items():
                        if clave != "archivo" and valor is not None:
                            if isinstance(valor, list):
                                metadatos_final[clave] = ", ".join(map(str, valor))
                            elif isinstance(valor, tipos_permitidos):
                                metadatos_final[clave] = valor
                            else:
                                metadatos_final[clave] = str(valor)
                
                    doc = Document(page_content=texto_limpio, metadata=metadatos_final)
                    docs_lista.append(doc) 
                    
        # LOGICA PARA WORD (DOCX)
        elif ruta_archivo.lower().endswith(".docx"):
            word_doc = docx.Document(ruta_archivo)
            texto_completo = "\n".join([p.text for p in word_doc.paragraphs if p.text.strip()])
            texto_limpio = limpiar_texto(texto_completo)
            
            if texto_limpio:
                metadatos_final = {
                    "fuente": archivo,
                    "doc_id": id_actual if id_actual else "unknown",
                    "tipo_documento": "Archivo Externo (Word)",
                    "pagina": 1, # Al no haber paginas claras, asignamos 1
                    "origen": origen_tipo 
                }
                
                tipos_permitidos = (str, int, float, bool)
                for clave, valor in diccionario_json.items():
                    if clave != "archivo" and valor is not None:
                        if isinstance(valor, list):
                            metadatos_final[clave] = ", ".join(map(str, valor))
                        elif isinstance(valor, tipos_permitidos):
                            metadatos_final[clave] = valor
                        else:
                            metadatos_final[clave] = str(valor)
                            
                doc = Document(page_content=texto_limpio, metadata=metadatos_final)
                docs_lista.append(doc)
            
        print(f"-> {archivo} leido correctamente en memoria.")
            
    except Exception as e:
        print(f"ERROR procesando {archivo}: {e}")
        return False, f"Error leyendo el archivo: {e}" 

    # 3. Chunking inteligente
    print("[3/4] Troceando el documento...")
    chunky = configurar_chunkeador(800, 100)
    chunks = chunky.split_documents(docs_lista) 
    print(f" -> Generados {len(chunks)} fragmentos.")

    # 4. Guardado en ChromaDB
    print("[4/4] Generando embeddings y guardando en BD...")
    try:
        ids = [generar_id_hash(chunk.page_content, chunk.metadata) for chunk in chunks]
        vector_db.add_documents(documents=chunks, ids=ids)
        return True, f"Exito: Se han anadido {len(chunks)} fragmentos a la base de datos."
    except Exception as e:
        for chunk_individual in chunks:
                try:
                    chunk_id = generar_id_hash(chunk_individual.page_content, chunk_individual.metadata)
                    vector_db.add_documents(documents=[chunk_individual], ids=[chunk_id])
                except Exception as ex:
                    print(f"\n[DESCARTADO] Exceso de tokens en ID {chunk_individual.metadata.get('doc_id')}")
                    continue
        return True, "Ingesta completada con descartes."